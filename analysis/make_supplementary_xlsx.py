"""Build Supplementary Data Sheet 1 (Tables S1-S7) from results/tables. Run after make_tables_figures.py."""
from decimal import Decimal, ROUND_HALF_EVEN
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = Path(__file__).resolve().parents[1]; T = ROOT / "results/tables"
out = ROOT / "results/Supplementary_Data_Sheet_1.xlsx"
LAB = {"no_treatment": "Untreated", "treat_hip": "Hip", "treat_knee": "Knee", "treat_ankle": "Ankle",
       "treat_hip_knee": "Hip + knee", "treat_hip_ankle": "Hip + ankle", "treat_knee_ankle": "Knee + ankle",
       "treat_all": "All three (reference)"}

def pct(df, cols):
    for c in cols:
        df[c] = (df[c] * 100).map(lambda v: v if pd.isna(v) else float(Decimal(repr(round(float(v), 9))).quantize(Decimal("0.1"), rounding=ROUND_HALF_EVEN)))
    return df

def tests(f):
    d = pd.read_csv(T / f)
    d["condition"] = d["condition"].map(LAB)
    d = pct(d, ["p_fall", "p_untreated", "delta_vs_untreated", "delta_ci_lo", "delta_ci_hi", "delta_vs_reference", "delta_ref_ci_hi"])
    d["wilcoxon_p_uncorrected"] = d["wilcoxon_p_uncorrected"].round(4)
    return d.rename(columns={"p_fall": "fall probability (%)", "p_untreated": "untreated (%)",
        "delta_vs_untreated": "difference vs untreated (pp)", "delta_ci_lo": "95% CI low (pp)", "delta_ci_hi": "95% CI high (pp)",
        "effect": "effect (CI excludes 0)", "seeds_improved": "policies improved", "seeds_worsened": "policies worsened",
        "wilcoxon_p_uncorrected": "Wilcoxon P (reference only)", "delta_vs_reference": "difference vs reference (pp)",
        "delta_ref_ci_hi": "95% CI high vs reference (pp)", "restored": "restored (CI high <= 5 pp)"})

s1 = pct(pd.read_csv(T / "TableS1_single_joint_dose_symmetric.csv"), ["mean", "std"]).rename(
    columns={"target": "perturbed joint(s)", "damping_scale": "damping scale (x)", "speed": "speed (m/s)",
             "mean": "fall probability mean (%)", "std": "SD across policies (%)"})
s2 = pct(pd.read_csv(T / "TableS2_dose_directional.csv"), ["mean", "std"]).rename(
    columns={"joint_set": "perturbed joint(s)", "kd": "resistance gain b (N m s/rad)", "speed": "speed (m/s)",
             "mean": "fall probability mean (%)", "std": "SD across policies (%)"})
s5 = pd.read_csv(T / "TableS5_kinematics_symmetric_all.csv"); s5["condition"] = s5["condition"].map(LAB)
s6 = pd.read_csv(T / "TableS6_kinematics_directional_all.csv"); s6["condition"] = s6["condition"].map(LAB)
s7 = pd.read_csv(T / "TableS7_reproducibility.csv").rename(columns={"kd": "resistance gain b (N m s/rad)"})
sheets = {
    "README": pd.DataFrame({"Sheet": ["Table S1", "Table S2", "Table S3", "Table S4", "Table S5", "Table S6", "Table S7", "Notes"],
        "Content": [
            "Single-joint dose-response, contracture-like model (actuator damping scale), all speeds and sides; scale 1 = unperturbed. 8 policies x 10 rollouts per cell.",
            "Dose-response, direction-selective velocity-dependent model (ankle alone or hip+knee+ankle), all speeds and sides.",
            "Relief effects, contracture-like model (damping x10,000 at hip pitch, knee, ankle pitch): difference vs untreated with hierarchical-bootstrap 95% CI, policies improved/worsened, Wilcoxon P (reference only), difference vs unperturbed reference and restoration.",
            "Relief effects, direction-selective model (right b = 40, left b = 50 N m s/rad).",
            "Kinematic and actuator measures (means over included rollouts) for every condition, side and speed, contracture-like model. time_to_fall_s mean time to fall over rollouts that fell (s); rom_* range of motion (deg); tauabs_* mean |actuator torque| (N m); taupk_* peak |actuator torque| (N m); sat_* fraction of time at the torque limit (|torque| >= 98% of limit); taudir_* mean |direction-selective torque| (N m); actstd_* SD of policy action; foot_* swing-peak measures (m, counts; main >= 0.05 m, secondary < 0.05 m); trunk_roll_std_deg, trunk_pitch_std_deg SD of pelvis roll and pitch (deg); trunk_tilt_max_deg maximum pelvis tilt (deg); base_vx_mean mean forward velocity (m/s); vx_tracking_err |mean forward velocity - commanded speed| (m/s).",
            "As Table S5, direction-selective model.",
            "Reproducibility: the untreated direction-selective three-joint condition run twice with identical configuration (gain sweep vs relief experiment).",
            "pp = percentage points. Kinematic and actuator measures are computed from 0.5 s to episode end or first fall; rollouts with fewer than 50 valid steps in this window are excluded from these measures (they remain in fall probability). Raw per-rollout data: results/raw/*.csv in the code repository."]}),
    "Table S1": s1, "Table S2": s2,
    "Table S3": tests("TableS3_symmetric_relief_tests.csv"), "Table S4": tests("TableS4_directional_relief_tests.csv"),
    "Table S5": s5.round(3), "Table S6": s6.round(3), "Table S7": s7.round(4)}
with pd.ExcelWriter(out, engine="openpyxl") as w:
    for name, df in sheets.items():
        df.to_excel(w, sheet_name=name, index=False)
wb = load_workbook(out)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=9, bold=(c.row == 1))
            if c.row == 1:
                c.fill = PatternFill("solid", fgColor="E8E8E8"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in ws.columns:
        L = max(len(str(c.value)) if c.value is not None else 0 for c in col[1:50]) if ws.max_row > 1 else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(10, L + 2), 22)
    ws.freeze_panes = "A2"
wb["README"].column_dimensions["B"].width = 120
for c in wb["README"]["B"]:
    c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(out); print("saved", out)
